"""Build the Module 3 industry transformation deck (10 slides, 16:9) with python-pptx.

Usage:  python build_deck.py
Output: Bandari_Ruthvik_Nath_Module3_AI_Industry_Transformation.pptx (next to this script)

The chart and all diagrams are native PowerPoint objects, so the deck stays sharp and editable.
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "Bandari_Ruthvik_Nath_Module3_AI_Industry_Transformation.pptx"

# ---- palette ----------------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x1B, 0x2D)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x1F, 0x5F, 0xAD)
GREY = RGBColor(0x9A, 0xA3, 0xAF)
RED = RGBColor(0xC0, 0x39, 0x2B)
TEAL = RGBColor(0x1E, 0x84, 0x6F)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
MARGIN = 0.6


# ---- helpers ----------------------------------------------------------------------------------
def rich(paragraph, text, size, color, bold=False, italic=False):
    """Add runs to a paragraph; **bold** and *italic* markers inside `text` are honoured."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        run = paragraph.add_run()
        is_b, is_i = part.startswith("**"), part.startswith("*") and not part.startswith("**")
        run.text = part.strip("*") if (is_b or is_i) else part
        f = run.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), color
        f.bold = bold or is_b
        f.italic = italic or is_i


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=6, italic=False):
    """Text box. `paras` is a str or list of str (one paragraph each)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p_text in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        rich(p, p_text, size, color, bold, italic)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, after=10, marker_color=ACCENT):
    """Bulleted list with a coloured square marker drawn as a real bullet character."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        p.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.28))))
        pPr.set("indent", str(-Emu(Inches(0.28))))
        bu_clr = pPr.makeelement(qn("a:buClr"), {})
        bu_clr.append(bu_clr.makeelement(qn("a:srgbClr"), {"val": str(marker_color)}))
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        rich(p, item, size, color)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def line(slide, x1, y1, x2, y2, color=RULE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def header(slide, title, kicker):
    rect(slide, 0, 0, 0.18, SH, ACCENT)
    text(slide, MARGIN, 0.42, SW - 2 * MARGIN, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True)
    text(slide, MARGIN, 0.72, SW - 2 * MARGIN, 0.8, title, size=30, color=NAVY, bold=True, spacing=1.0)


def footer(slide, source, n):
    line(slide, MARGIN, SH - 0.55, SW - MARGIN, SH - 0.55, RULE, 0.75)
    text(slide, MARGIN, SH - 0.45, SW - 2 * MARGIN - 0.6, 0.3, source, size=10, color=MUTED)
    text(slide, SW - MARGIN - 0.5, SH - 0.45, 0.5, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = " ".join(body.split())


def card(slide, x, y, w, h, title, tag, rows, color=ACCENT):
    """Organization card: coloured top bar, name, technique tag, then labelled rows."""
    rect(slide, x, y, w, h, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    rect(slide, x, y, w, 0.09, color)
    text(slide, x + 0.25, y + 0.25, w - 0.5, 0.4, title, size=20, color=NAVY, bold=True, after=0)
    text(slide, x + 0.25, y + 0.68, w - 0.5, 0.35, tag, size=12, color=color, bold=True, after=0)
    paras = []
    for label, body in rows:
        paras.append(f"**{label}**")
        paras.append(body)
    box = text(slide, x + 0.25, y + 1.1, w - 0.5, h - 1.25, paras, size=16, color=INK, after=2)
    for i, p in enumerate(box.text_frame.paragraphs):       # gap before each label after the first
        if i and i % 2 == 0:
            p.space_before = Pt(8)
    return box


def bar_chart(slide, x, y, w, h, title, categories, values, colors, number_format, axis_max):
    data = CategoryChartData()
    data.categories = categories
    data.add_series(title, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold = Pt(14), True
    tp.runs[0].font.color.rgb, tp.runs[0].font.name = NAVY, FONT
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = number_format, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.color.rgb, dl.font.name = Pt(16), True, INK, FONT
    for i, c in enumerate(colors):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = axis_max, 0
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.color.rgb, ca.tick_labels.font.name = Pt(12), INK, FONT
    ca.format.line.color.rgb = RULE
    return gf


# ---- slide 1: title ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, MARGIN, 1.55, 0.9, 0.08, TEAL)
text(s, MARGIN, 1.8, 12.0, 1.7, ["From Signatures to Signals:", "How AI Is Transforming Cyber Defense"],
     size=44, color=WHITE, bold=True, spacing=1.0, after=0)
text(s, MARGIN, 3.6, 12.0, 1.0, "A strategic briefing for leadership on AI-driven threat detection and threat intelligence",
     size=22, color=RGBColor(0xC9, 0xD3, 0xE0), spacing=1.1)
text(s, MARGIN, 5.25, 12.0, 1.4, [
    "Ruthvik Nath Bandari",
    "AAI6680: AI for Cybersecurity · Northeastern University · Mimoza Dimodugno, PhD",
    "October 2, 2026",
], size=16, color=RGBColor(0xC9, 0xD3, 0xE0), after=4)
notes(s, """This briefing answers one question for leadership: how is AI changing security operations at
leading organizations, and should we follow? I look at two areas: AI for threat detection, and AI for
real-time threat intelligence, and close with a recommendation and what this research means for me.""")

# ---- slide 2: executive summary ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "AI is now essential to cyber defense, but it needs oversight", "Executive summary")
bullets(s, MARGIN, 1.85, 7.3, 4.6, [
    "**Threats outpace signatures:** 82% of CrowdStrike's 2025 detections were malware-free.",
    "**AI pays off:** extensive security AI was linked to **$1.9M** lower breach costs and **80** fewer days.",
    "**Intelligence turns proactive:** Google predicted a SQLite exploit and cut it off first.",
    "**Attackers use AI too:** one actor compromised 600+ FortiGate devices using commercial AI.",
], size=18, after=16)
rect(s, 8.35, 1.85, 4.4, 3.6, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
rect(s, 8.35, 1.85, 4.4, 0.09, TEAL)
text(s, 8.65, 2.15, 3.8, 0.3, "RECOMMENDATION", size=12, color=TEAL, bold=True, after=0)
text(s, 8.65, 2.55, 3.8, 3.7, [
    "Adopt AI detection and triage in phases.",
    "Keep analysts approving high-impact actions.",
    "Govern every AI tool (NIST CSF 2.0, AI RMF).",
], size=18, color=INK, after=14)
footer(s, "Sources: CrowdStrike (2026b); IBM (2025); Walker (2025); Krebs (2026); National Institute of Standards and Technology [NIST] (2023, 2024).", 2)
notes(s, """Four findings. First, attacks increasingly involve no malware: CrowdStrike reports that 82% of its
2025 detections were malware-free, which leaves signature tools nothing to match. Second, the return is
measurable: IBM's 2025 Cost of a Data Breach study found that organizations using security AI and automation
extensively had breaches that cost 1.9 million dollars less and lasted 80 fewer days than organizations using
none. IBM presents this as a correlation, and its 2026 edition reports a similar 1.93 million dollar gap.
Third, threat intelligence is becoming predictive: Google says that by combining its threat intelligence with
its Big Sleep AI agent, it predicted a SQLite vulnerability was about to be exploited and cut it off first.
Fourth, attackers adopt AI as fast as defenders, as Brian Krebs reported. My recommendation is phased
adoption, with analysts approving high-impact actions and governance aligned with NIST CSF 2.0 and the NIST
AI Risk Management Framework.""")

# ---- slide 3: Step 1, three organizations ------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Three leaders, three AI techniques", "Step 1 · AI-driven threat detection · Q1")
cw, cy, ch_ = 3.9, 1.9, 4.1
card(s, MARGIN, cy, cw, ch_, "Microsoft", "GENERATIVE AI (LLM) TRIAGE", [
    ("Use", "LLM agents help analysts triage phishing and incidents."),
    ("Result", "Trial: **77% better** verdict accuracy. Live SOCs: 30% shorter resolution time. Both are Microsoft studies."),
], ACCENT)
card(s, MARGIN + cw + 0.32, cy, cw, ch_, "CrowdStrike", "BEHAVIORAL MACHINE LEARNING", [
    ("Use", "Models attacker behavior across endpoint, identity and cloud."),
    ("Result", "MITRE 2025: **100% detection, zero false positives** (vendor's reading of an independent test)."),
], TEAL)
card(s, MARGIN + 2 * (cw + 0.32), cy, cw, ch_, "Palo Alto Networks", "INLINE DEEP LEARNING ON FILES", [
    ("Use", "Classifies files in real time, including files no signature covers yet."),
    ("Result", "Stops **99%+ of unknown malware** (vendor-reported)."),
], RED)
footer(s, "Sources: Bono (2025); Bono et al. (2024); CrowdStrike (2025b); Palo Alto Networks (n.d.).", 3)
notes(s, """Which AI techniques are most effective? I chose three organizations that use different ones.
Microsoft applies generative AI to triage. In a randomized controlled trial by Microsoft researchers, analysts
using a phishing-triage agent achieved a 77% improvement in verdict accuracy. A separate observational study of
live security operations centers linked Security Copilot adoption to a 30% reduction in mean time to resolve
incidents, although its authors note it cannot prove causation. Both studies are Microsoft's own. CrowdStrike
relies on behavioral machine learning across endpoint, identity, and cloud; in MITRE's 2025 ATT&CK enterprise
evaluation, which emulated Scattered Spider and Mustang Panda, CrowdStrike reports 100% detection with no false
positives. MITRE runs the test but does not rank vendors, so that headline is CrowdStrike's interpretation. Palo
Alto Networks uses inline machine learning and deep learning models that classify files from their structure,
covering files no signature exists for yet; its claim of stopping over 99% of unknown malware is marketing data
that it pairs with faster signature generation, not an independent audit.""")

# ---- slide 4: AI vs signatures + ROI chart ----------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "AI beats signatures on what is new, and the return is measurable", "Step 1 · Effectiveness and ROI · Q2")
bullets(s, MARGIN, 1.85, 4.6, 4.6, [
    "**Signatures miss the new:** 23% of detected malware evaded them in late 2025 (vendor telemetry).",
    "**Detection improved:** a Fortune 500 AI pilot found threats 26–36% faster, with false positives down 16 points.",
    "**The savings held in 2026:** $1.93M lower cost per breach.",
], size=16, after=16)
bar_chart(s, 5.55, 1.75, 3.55, 4.35, "Breach cost (USD millions)",
          ["No security AI", "Extensive AI"], (5.52, 3.62), (GREY, ACCENT), '"$"0.00', 6.6)
bar_chart(s, 9.2, 1.75, 3.55, 4.35, "Days to identify and contain",
          ["No security AI", "Extensive AI"], (284, 204), (GREY, ACCENT), "0", 340)
text(s, 5.55, 6.15, 7.2, 0.6, "**−$1.90M and −80 days** with extensive security AI vs none (IBM, 2025). A correlation, not proof of cause.",
     size=13, color=MUTED, align=PP_ALIGN.CENTER)
footer(s, "Sources: WatchGuard Technologies (2026); Bracken (2026); IBM (2025, 2026).", 4)
notes(s, """How much have organizations improved detection with AI? Signature-based detection compares files to
known bad patterns, so it fails on anything new. WatchGuard's vendor telemetry for the second half of 2025 found
that 23% of detected malware evaded signature-based detection, while new unique malware rose 1,548% from the
third to the fourth quarter. The clearest field result I found is a six-month AI SOC pilot at a Fortune 500
manufacturer, presented at RSAC 2026 and reported by Dark Reading: mean time to discovery improved 26 to 36%,
mean time to respond improved 22%, and false positives fell 16 points, although the AI also introduced some new
false alerts. These are the practitioner's own measurements. The chart shows return on investment. In IBM's
2025 study, run by the Ponemon Institute, breaches at organizations with no security AI cost an average of 5.52
million dollars and took 284 days to identify and contain, against 3.62 million dollars and 204 days with
extensive use. The 2026 edition reports a 1.93 million dollar gap. These are cross-organization correlations,
not controlled experiments.""")

# ---- slide 5: what works and challenges --------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "What works best, and what gets in the way", "Step 1 · Techniques and challenges · Q1 · Q3")
rect(s, MARGIN, 1.85, 5.7, 3.4, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
text(s, MARGIN + 0.3, 2.1, 5.1, 0.3, "MOST EFFECTIVE TECHNIQUES", size=12, color=TEAL, bold=True, after=0)
bullets(s, MARGIN + 0.3, 2.55, 5.1, 3.8, [
    "**Behavioral ML** for malware-free attacks",
    "**Deep learning** for unseen malware files",
    "**LLM assistants** for faster triage",
    "Best results: all three, plus **human review**",
], size=18, after=14, marker_color=TEAL)
text(s, 6.75, 1.9, 6.0, 0.3, "IMPLEMENTATION CHALLENGES", size=12, color=RED, bold=True, after=0)
bullets(s, 6.75, 2.35, 6.0, 4.1, [
    "**Noise:** 55% face too many false positives",
    "**Trust:** only 11% fully trust AI for critical tasks",
    "**Skills:** 59% report critical skill needs",
    "**AI as a target:** 97% of firms with AI breaches lacked AI access controls",
], size=18, after=14, marker_color=RED)
footer(s, "Sources: Cisco (2025a); CrowdStrike (2026b); ISC2 (2025); IBM (2025).", 5)
notes(s, """Which techniques work best? The evidence points to a combination. Behavioral machine learning and
anomaly detection matter most because 82% of CrowdStrike's detections were malware-free. Deep learning on file
structure covers new malware that signatures miss. Large language model assistants mainly add speed and
consistency to triage. What challenges do organizations face? Cisco's 2025 State of Security survey of 2,058
security leaders found 55% handle too many false positives and only 11% trust AI completely for
mission-critical tasks. In ISC2's 2025 workforce study, 59% of respondents reported critical or significant
skill needs. AI also becomes a target: IBM found that 97% of organizations reporting breaches of AI models or
applications lacked proper AI access controls. My own Module 3 lab showed a related issue: a model can hold the
right signal and still miss an attack because of how it scores anomalies, so tuning and validation matter.""")

# ---- slide 6: Step 2, vendors -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "AI threat intelligence at machine speed", "Step 2 · Real-time threat intelligence · Q5")
text(s, MARGIN, 1.65, SW - 2 * MARGIN, 0.4,
     "Why speed matters: in 2025, average eCrime breakout time fell to **29 minutes**.", size=16, color=INK)
cy, ch_ = 2.35, 3.95
card(s, MARGIN, cy, cw, ch_, "CrowdStrike", "CHARLOTTE AI", [
    ("SOC tool", "Agentic alert triage, trained on millions of analyst decisions."),
    ("Result", "**98%+ triage accuracy** and about 40 analyst hours saved a week (vendor)."),
], ACCENT)
card(s, MARGIN + cw + 0.32, cy, cw, ch_, "SentinelOne", "PURPLE AI", [
    ("SOC tool", "Agentic investigation, opened to all customers in 2026."),
    ("Result", "MITRE 2024: all 16 attack steps detected, **88% fewer alerts** than the median (vendor)."),
], TEAL)
card(s, MARGIN + 2 * (cw + 0.32), cy, cw, ch_, "Splunk (Cisco)", "AI IN THE SIEM", [
    ("SOC tool", "Enterprise Security 8.2, with Triage and Malware Reversal agents announced alongside (Triage due 2026)."),
    ("Industry survey", "59% of security leaders say AI raised their efficiency."),
], RED)
footer(s, "Sources: CrowdStrike (2025a, 2026a); SentinelOne (2024, 2026); Cisco (2025a, 2025b).", 6)
notes(s, """Speed explains why threat intelligence has to be automated: CrowdStrike's 2026 Global Threat Report
measured average eCrime breakout time at 29 minutes, with the fastest at 27 seconds. Which AI tools are SOC teams
using? CrowdStrike's Charlotte AI Detection Triage was trained on millions of real triage decisions by its
managed-detection analysts; CrowdStrike reports over 98% triage accuracy and estimates more than 40 hours of
manual work saved per week. SentinelOne's Purple AI opened agentic investigation to all customers in June 2026,
and in the 2024 MITRE evaluation SentinelOne reports detecting all 16 attack steps with 88% fewer alerts than the
median vendor. Splunk, now part of Cisco, brings AI into the SIEM: alongside Enterprise Security 8.2, Cisco
announced a Triage Agent that prioritizes and explains alerts, scheduled for 2026, and a Malware Reversal Agent
that explains malicious scripts and extracts indicators. Separately, in Cisco's 2025 industry survey, 59% of
security leaders said AI improved their efficiency. All product figures are vendor-reported.""")

# ---- slide 7: reactive vs proactive -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "From reactive to proactive threat intelligence", "Step 2 · Predictive analytics and hunting · Q4 · Q6")
rows = [
    ("", "Reactive (traditional)", "Proactive (AI-driven)"),
    ("Trigger", "Known IOCs and signatures", "Behavior and predicted exploits"),
    ("Timing", "After compromise", "Before or during breakout"),
    ("Hunting", "Manual expert queries", "Agentic investigation"),
    ("Alerts", "40% never investigated", "AI agents triage and explain"),
]
tx, ty, cws, rh = MARGIN, 1.85, (1.5, 2.6, 3.2), 0.62
for r, row in enumerate(rows):
    x = tx
    for c, cell in enumerate(row):
        if r == 0 and c > 0:
            rect(s, x, ty, cws[c] - 0.06, rh - 0.08, GREY if c == 1 else TEAL)
            text(s, x + 0.15, ty + 0.13, cws[c] - 0.3, rh, cell, size=15, color=WHITE, bold=True)
        elif r > 0:
            text(s, x + (0 if c == 0 else 0.15), ty + r * rh + 0.12, cws[c] - 0.3, rh, cell, size=15,
                 color=MUTED if c == 0 else INK, bold=(c == 0))
        x += cws[c]
    if r > 0:
        line(s, tx, ty + r * rh + rh, tx + sum(cws), ty + r * rh + rh, RULE, 0.75)
ex = MARGIN + sum(cws) + 0.35
text(s, ex, 1.85, SW - MARGIN - ex, 0.3, "EVIDENCE", size=12, color=ACCENT, bold=True, after=0)
bullets(s, ex, 2.25, SW - MARGIN - ex, 4.3, [
    "**Zero-day stopped:** Google's threat intelligence plus its Big Sleep AI predicted a SQLite exploit (CVE-2025-6965) and cut it off.",
    "**Early warning:** Darktrace flagged activity later traced to an Ivanti zero-day, 11 days before disclosure (vendor).",
    "**How well does AI predict?** Promising, but I found no independent study.",
], size=15, after=12)
footer(s, "Sources: Townsend (2025); Walker (2025); Turnbull (2025).", 7)
notes(s, """How is AI changing threat intelligence from reactive to proactive? Traditional intelligence waits for
indicators of compromise, such as known hashes and IP addresses, and responds after the damage. Reactive teams
are already overwhelmed: a vendor-sponsored survey of 282 security leaders, reported by SecurityWeek, found 40%
of alerts are never investigated. AI-driven intelligence uses behavior and context, so it can act earlier, and
AI agents from CrowdStrike, SentinelOne and Splunk now triage and explain alerts. Two documented cases show the
shift. In July 2025, Google said that combining its threat intelligence with its Big Sleep agent let it predict
that a SQLite vulnerability, CVE-2025-6965, was about to be exploited and cut it off first. Darktrace reports
that its models flagged suspicious activity in December 2024 that was later traced to exploitation of an Ivanti
zero-day, 11 days before Ivanti disclosed it. How effective is AI at predicting attacks overall? Promising but
unproven: every figure I found comes from a vendor, and I found no independent study of prediction accuracy.""")

# ---- slide 8: incident, CIA, recommendation, reflection -----------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Recommendation and reflection", "Case, CIA triad and strategy")
rect(s, MARGIN, 1.85, 6.15, 3.85, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
text(s, MARGIN + 0.3, 2.05, 5.6, 0.3, "CASE: AI-ASSISTED FORTIGATE CAMPAIGN (KREBS, 2026)", size=12, color=RED, bold=True, after=0)
text(s, MARGIN + 0.3, 2.42, 5.6, 1.6, [
    "**C:** network maps and credentials exposed",
    "**I:** firewall devices compromised by the attacker",
    "**A:** at risk if those firewalls fail (my inference)",
], size=14.5, color=INK, after=3)
text(s, MARGIN + 0.3, 3.75, 5.6, 0.3, "RECOMMENDATION", size=12, color=TEAL, bold=True, after=0)
bullets(s, MARGIN + 0.3, 4.12, 5.6, 2.4, [
    "**Protect:** MFA and no exposed management ports",
    "**Detect and respond:** behavioral AI with analyst approval",
    "**Govern:** access controls on every AI tool (AI RMF)",
], size=14.5, after=6, marker_color=TEAL)
text(s, 7.15, 1.9, 5.6, 0.3, "WHAT THIS MEANS FOR ME", size=12, color=ACCENT, bold=True, after=0)
text(s, 7.15, 2.35, 5.6, 4.2, [
    "My Module 1 and Module 3 detectors scored well, yet both missed attacks the headline numbers hid. I now read vendor claims the same way.",
    "As a future security engineer, I want to build behavioral detectors with measured, explained limits, and a named analyst accountable for what they automate.",
], size=16, color=INK, after=16)
footer(s, "Sources: Krebs (2026); NIST (2023, 2024). Reflection draws on my Module 1 and Module 3 labs.", 8)
notes(s, """To ground the recommendation, take the incident Brian Krebs reported in March 2026, citing Amazon AWS:
a low-skilled, Russian-speaking actor used commercial generative AI services to compromise more than 600 FortiGate
security appliances in at least 55 countries, by finding exposed management ports and weak single-factor
credentials. In CIA terms, confidentiality was lost because the actor obtained internal network topology and
confirmed credentials; integrity was lost because the security devices themselves were compromised; and
availability was at risk, since a compromised firewall can be disabled, which is my inference rather than a
reported outcome. AWS noted the actor simply moved on when it met hardened environments. So my recommendation,
aligned with NIST CSF 2.0, is to protect first with multi-factor authentication and no internet-exposed
management interfaces; then detect and respond with behavioral AI and AI-assisted triage, keeping analysts
approving high-impact actions; and govern every AI tool with access controls following the NIST AI Risk
Management Framework. Personally, my Module 1 Gaussian mixture detector missed SSH brute force, and my Module 3
autoencoder missed most password guessing because averaging its error hid a clear signal. Both scored well
overall. That is why I read vendor claims critically, and why I want to build behavioral detectors whose limits
are measured and explained.""")

# ---- slides 9-10: references ---------------------------------------------------------------------
REFS = [
    "Bono, J. (2025). *Randomized controlled trials for phishing triage agent* (arXiv:2511.13860). arXiv. https://doi.org/10.48550/arXiv.2511.13860",
    "Bono, J., Grana, J., & Xu, A. (2024). *Generative AI and security operations center productivity: Evidence from live operations* (arXiv:2411.03116). arXiv. https://doi.org/10.48550/arXiv.2411.03116",
    "Bracken, B. (2026, March 23). AI in the SOC: What could go wrong? *Dark Reading*. https://www.darkreading.com/cybersecurity-operations/ai-soc-go-wrong",
    "Cisco. (2025a, May 20). *Global state of security report reveals critical need for connected security operations* [Press release]. https://investor.cisco.com/news/news-details/2025/Global-State-of-Security-Report-Reveals-Critical-Need-for-Connected-Security-Operations/default.aspx",
    "Cisco. (2025b, September 9). *Cisco elevates the SOC with agentic AI for faster threat response and reduced complexity* [Press release]. https://newsroom.cisco.com/c/r/newsroom/en/us/a/y2025/m09/cisco-elevates-the-soc-with-agentic-ai-for-faster-threat-response-and-reduced-complexity.html",
    "CrowdStrike. (2025a, February 13). *CrowdStrike unveils Charlotte AI Detection Triage for faster SOC triage* [Press release]. https://www.crowdstrike.com/en-us/press-releases/crowdstrike-delivers-next-breakthrough-in-ai-powered-agentic-cybersecurity-with-charlotte-ai-detection-triage/",
    "CrowdStrike. (2025b, December 10). *CrowdStrike achieves 100% in 2025 MITRE ATT&CK Enterprise Evaluation*. https://www.crowdstrike.com/en-us/blog/crowdstrike-achieves-100-percent-2025-mitre-attack-enterprise-evaluation/",
    "CrowdStrike. (2026a, February 24). *2026 CrowdStrike global threat report: AI accelerated adversaries* [Press release]. https://www.crowdstrike.com/en-us/press-releases/2026-crowdstrike-global-threat-report/",
    "CrowdStrike. (2026b, February 24). *CrowdStrike 2026 global threat report: Evasive adversary wields AI*. https://www.crowdstrike.com/en-us/blog/crowdstrike-2026-global-threat-report-findings/",
    "IBM. (2025). *Cost of a data breach report 2025: The AI oversight gap*. https://www-api.ibm.com/adobe/assets/urn:aaid:aem:607b9590-38e0-4c91-b433-aa8a17f5b5e8/original/as/cost-of-a-data-breach-2025-full-report.pdf",
    "IBM. (2026). *Cost of a data breach report 2026*. https://www.ibm.com/reports/data-breach",
    "ISC2. (2025, December 4). *ISC2 study finds cybersecurity budget constraints remain, but do not worsen, while skill needs grow*. https://www.isc2.org/insights/2025/12/isc2-publishes-2025-cybersecurity-workforce-study",
    "Krebs, B. (2026, March 8). How AI assistants are moving the security goalposts. *Krebs on Security*. https://krebsonsecurity.com/2026/03/how-ai-assistants-are-moving-the-security-goalposts/",
    "National Institute of Standards and Technology. (2023). *Artificial intelligence risk management framework (AI RMF 1.0)* (NIST AI 100-1). https://doi.org/10.6028/NIST.AI.100-1",
    "National Institute of Standards and Technology. (2024). *The NIST cybersecurity framework (CSF) 2.0* (NIST CSWP 29). https://doi.org/10.6028/NIST.CSWP.29",
    "Palo Alto Networks. (n.d.). *Advanced WildFire*. Retrieved October 2, 2026, from https://www.paloaltonetworks.com/network-security/advanced-wildfire",
    "SentinelOne. (2024, December 11). *SentinelOne sets the standard with 100% detection and 88% fewer alerts than median across all vendors evaluated in the 2024 MITRE ATT&CK Evaluations: Enterprise* [Press release]. https://www.sentinelone.com/press/sentinelone-sets-the-standard-with-100-detection-and-88-fewer-alerts-than-median-across-all-vendors-evaluated-in-the-2024-mitre-attck-evaluations-enterprise/",
    "SentinelOne. (2026, June 17). *SentinelOne opens Purple AI agentic investigation to all customers, bringing frontier AI directly into the SOC* [Press release]. https://www.sentinelone.com/press/sentinelone-opens-purple-ai-agentic-investigation-to-all-customers-bringing-frontier-ai-directly-into-the-soc/",
    "Townsend, K. (2025, September 11). AI emerges as the hope—and risk—for overloaded SOCs. *SecurityWeek*. https://www.securityweek.com/ai-emerges-as-the-hope-and-risk-for-overloaded-socs/",
    "Turnbull, H. (2025, March 7). *Darktrace's early detection of the latest Ivanti exploits*. Darktrace. https://www.darktrace.com/blog/darktraces-early-detection-of-the-latest-ivanti-exploits",
    "Walker, K. (2025, July 15). *Google's latest AI security announcements*. Google. https://blog.google/technology/safety-security/cybersecurity-updates-summer-2025/",
    "WatchGuard Technologies. (2026, February 19). *Over 1500% increase in new, unique malware highlights growing security complexity, according to WatchGuard biannual threat report* [Press release]. https://www.watchguard.com/wgrd-news/press-releases/over-1500-increase-new-unique-malware-highlights-growing-security",
]
half = 11
for page, chunk in enumerate((REFS[:half], REFS[half:])):
    s = prs.slides.add_slide(BLANK)
    header(s, "References", f"APA 7th edition · {page + 1} of 2")
    box = text(s, MARGIN, 1.7, SW - 2 * MARGIN, 5.2, chunk, size=10.5, color=INK, after=4, spacing=1.0)
    for p in box.text_frame.paragraphs:        # APA hanging indent
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.45))))
        pPr.set("indent", str(-Emu(Inches(0.45))))
    footer(s, "Deck structure: title, seven content slides, two reference slides. Vendor figures are labelled.", 9 + page)
    notes(s, "References in APA 7th edition. Vendor figures are labelled where they appear.")

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides")
