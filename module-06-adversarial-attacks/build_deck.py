"""Build the Module 6 adversarial attack analysis deck (16:9, seven slides) with python-pptx.

Usage:  python build_deck.py
Output: Bandari_Ruthvik_Nath_Module6_Adversarial_Attack_Analysis.pptx
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "Bandari_Ruthvik_Nath_Module6_Adversarial_Attack_Analysis.pptx"
DATE = "October 19, 2026"

# ---- palette ----------------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x1B, 0x2D)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x1F, 0x5F, 0xAD)
GREY = RGBColor(0x9A, 0xA3, 0xAF)
RED = RGBColor(0xC0, 0x39, 0x2B)
TEAL = RGBColor(0x1E, 0x84, 0x6F)
SOFT = RGBColor(0xC9, 0xD3, 0xE0)
GOLD = RGBColor(0xB7, 0x79, 0x1F)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
MARGIN = 0.6


# ---- helpers ----------------------------------------------------------------------------------
def rich(paragraph, text, size, color, bold=False, italic=False):
    """Add runs to a paragraph; **bold** and *italic* markers inside `text` are honoured."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        run = paragraph.add_run()
        is_b, is_i = part.startswith("**"), part.startswith("*") and not part.startswith("**")
        run.text = part.strip("*") if (is_b or is_i) else part
        f = run.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), color
        f.bold = bold or is_b
        f.italic = italic or is_i


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=6, italic=False):
    """Text box. `paras` is a str or list of str (one paragraph each)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p_text in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        rich(p, p_text, size, color, bold, italic)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, after=10, marker_color=ACCENT):
    """Bulleted list with a coloured square marker drawn as a real bullet character."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        p.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.28))))
        pPr.set("indent", str(-Emu(Inches(0.28))))
        bu_clr = pPr.makeelement(qn("a:buClr"), {})
        bu_clr.append(bu_clr.makeelement(qn("a:srgbClr"), {"val": str(marker_color)}))
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        rich(p, item, size, color)
    return box


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def line(slide, x1, y1, x2, y2, color=RULE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def header(slide, title, kicker):
    rect(slide, 0, 0, 0.18, SH, ACCENT)
    text(slide, MARGIN, 0.42, SW - 2 * MARGIN, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True)
    text(slide, MARGIN, 0.72, SW - 2 * MARGIN, 0.8, title, size=30, color=NAVY, bold=True, spacing=1.0)


def footer(slide, source, n):
    line(slide, MARGIN, SH - 0.55, SW - MARGIN, SH - 0.55, RULE, 0.75)
    text(slide, MARGIN, SH - 0.45, SW - 2 * MARGIN - 0.6, 0.3, source, size=10, color=MUTED)
    text(slide, SW - MARGIN - 0.5, SH - 0.45, 0.5, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = " ".join(body.split())


def stat(slide, x, y, w, value, label, color=ACCENT):
    rect(slide, x, y, w, 1.35, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(slide, x, y, 0.08, 1.35, color)
    text(slide, x + 0.3, y + 0.14, w - 0.45, 0.6, value, size=30, color=color, bold=True, after=0)
    text(slide, x + 0.3, y + 0.78, w - 0.45, 0.55, label, size=13, color=INK, after=0, spacing=1.0)




def case_card(slide, x, y, w, h, color, sector, title, rows):
    rect(slide, x, y, w, h, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    rect(slide, x, y, w, 0.09, color)
    text(slide, x + 0.25, y + 0.25, w - 0.5, 0.3, sector.upper(), size=11, color=color, bold=True, after=0)
    text(slide, x + 0.25, y + 0.55, w - 0.5, 0.75, title, size=18, color=NAVY, bold=True, after=0, spacing=1.0)
    paras = []
    for label, body in rows:
        paras.append(f"**{label}** {body}" if label else body)
    text(slide, x + 0.25, y + 1.35, w - 0.5, h - 1.5, paras, size=13, color=INK, after=7)


CH5, CH10 = "Parisi (2019, Ch. 5)", "Parisi (2019, Ch. 10)"

# ---- slide 1: title ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, MARGIN, 1.4, 0.9, 0.08, RED)
text(s, MARGIN, 1.65, 12.0, 1.7, ["When AI Is the Target:", "Real-World Adversarial Attacks and Defenses"],
     size=42, color=WHITE, bold=True, spacing=1.0, after=0)
text(s, MARGIN, 3.45, 12.0, 1.0, "One threat, adversarial manipulation of AI systems, traced through three real cases, "
     "three company defenses and the rules now taking shape", size=20, color=SOFT, spacing=1.1)
text(s, MARGIN, 4.85, 12.0, 1.4, ["Ruthvik Nath Bandari",
                                    "AAI6680: AI for Cybersecurity · Northeastern University · Mimoza Dimodugno, PhD", DATE],
     size=16, color=SOFT, after=4)
notes(s, """This presentation analyses one cyber threat: adversarial attacks, where an attacker manipulates the input,
the training data or the instructions of an AI system so that it makes a wrong decision. I look at three real cases, how
three major companies defend their AI, how the defenses compare, and how anomaly detection and AI governance help.""")

# ---- slide 2: the threat ------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "The threat: tricking AI into the wrong decision", "Threat and impact")
TECH = [("Evasion", "Altered input fools a trained model (a sign, an image)", ACCENT),
        ("Poisoning", "Corrupted training data or model components", TEAL),
        ("Prompt injection", "Hidden instructions hijack an AI assistant", RED),
        ("Deepfakes", "AI-faked faces and voices fool people and checks", GOLD)]
for i, (t_, d_, c_) in enumerate(TECH):
    x = MARGIN + i * 3.07
    rect(s, x, 1.85, 2.87, 1.55, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(s, x, 1.85, 2.87, 0.08, c_)
    text(s, x + 0.2, 2.05, 2.5, 0.4, t_, size=17, color=c_, bold=True, after=0)
    text(s, x + 0.2, 2.5, 2.5, 0.85, d_, size=13, color=INK, after=0)
stat(s, MARGIN, 3.75, 3.9, "US$25 million", "lost in one deepfake video call, Hong Kong, January 2024", RED)
stat(s, MARGIN + 4.12, 3.75, 3.9, "US$40 billion", "projected US fraud losses enabled by generative AI by 2027 (32% a year)", RED)
stat(s, MARGIN + 8.24, 3.75, 3.9, "40+ million", "vehicles use the camera that a strip of tape fooled", ACCENT)
text(s, MARGIN, 5.45, 12.1, 1.2, [
    "**Most exposed industries:** autonomous vehicles (safety), finance (money moved on a faked instruction) and "
    "healthcare (life-critical model decisions).",
], size=15, color=INK)
footer(s, "Sources: Lalchand et al. (2024); Povolny & Trivedi (2020); Parisi (2019, Ch. 10).", 2)
notes(s, """An adversarial attack makes an AI system decide wrongly on purpose. The four techniques in the brief are evasion,
where a crafted input fools a trained model; poisoning, where the training data or model is corrupted; prompt injection,
where hidden instructions take over an AI assistant; and deepfakes, where AI-generated faces and voices impersonate real
people. The impact is financial. In January 2024 one deepfake video call cost a Hong Kong firm 25 million US dollars, and
Deloitte projects that generative AI could push US fraud losses to 40 billion dollars by 2027, growing 32 percent a year.
The camera fooled by tape in the McAfee test is used in more than 40 million vehicles. Deepfake fraud is the fastest-growing
technique against people, and prompt injection is the main attack on AI assistants. The most exposed industries are autonomous
vehicles, finance and healthcare, because a single wrong decision there costs lives or money.""")

# ---- slide 3: real cases -------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Three real attacks that made headlines", "Step 1 · case studies")
cw = 3.95
case_card(s, MARGIN, 1.8, cw, 3.75, ACCENT, "Autonomous vehicles · evasion", "Tesla: tape on a speed sign", [
    ("How:", "5 cm of black tape on a 35 mph sign made the MobilEye camera read 85."),
    ("Impact:", "The car began accelerating toward 85 mph; the camera is in 40+ million vehicles."),
    ("Outcome:", "Newer cameras handle it; older cars were not updated."),
])
case_card(s, MARGIN + cw + 0.18, 1.8, cw, 3.75, RED, "Finance · deepfake", "Hong Kong: a fake CFO on video", [
    ("How:", "An employee joined a video call where the CFO and colleagues were all deepfakes."),
    ("Impact:", "**US$25 million** sent to the fraudsters."),
    ("Outcome:", "Money lost; the firm and its staff were tricked, not its IT systems."),
])
case_card(s, MARGIN + 2 * (cw + 0.18), 1.8, cw, 3.75, GOLD, "Enterprise AI · prompt injection", "EchoLeak: Microsoft 365 Copilot", [
    ("How:", "A crafted email could make Copilot leak data with zero clicks."),
    ("Impact:", "Rated **9.3 Critical** (CVSS)."),
    ("Outcome:", "Microsoft fixed it server-side; no exploitation was reported."),
])
text(s, MARGIN, 5.8, 12.1, 0.3, "WHY THEY MADE HEADLINES", size=12, color=RED, bold=True, after=0)
text(s, MARGIN, 6.15, 12.1, 0.6, "None needed a break-in: each fooled a trusted AI, or a person, with input crafted for it.",
     size=15, color=INK)
footer(s, "Sources: Povolny & Trivedi (2020); Lalchand et al. (2024); Microsoft Security Response Center (2025).", 3)
notes(s, """Three cases show the threat across industries. In 2020, McAfee researchers put about two inches of black tape on a
35 mile-per-hour sign; the MobilEye camera on a 2016 Tesla read it as 85, and with cruise control on the car began
accelerating toward 85 before the researchers braked. That camera is in more than 40 million vehicles. In January 2024, an
employee of a Hong Kong firm joined a video call in which the chief financial officer and other colleagues were deepfakes,
and sent 25 million US dollars to the fraudsters. In June 2025, Microsoft disclosed
EchoLeak, a critical prompt-injection flaw rated 9.3 in which a crafted email could make Microsoft 365 Copilot leak data
without the user clicking anything. Microsoft fixed it on the server side, and no exploitation was reported, which makes it an
example of an attack stopped before it caused damage.""")

# ---- slide 4: corporate defenses ----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "How Microsoft, Google and Meta defend their AI", "Step 2 · corporate defenses")
case_card(s, MARGIN, 1.8, cw, 3.55, ACCENT, "Microsoft · red teaming", "Attack it before attackers do", [
    ("", "Red-teamed **100+** generative AI products."),
    ("", "Lessons: attacks need no gradients, humans stay crucial, security is never finished."),
])
case_card(s, MARGIN + cw + 0.18, 1.8, cw, 3.55, TEAL, "Google · adversarial training", "Layered defense for Gemini", [
    ("", "Training on adversarial data “significantly enhanced” resistance to prompt injection."),
    ("", "Plus ML classifiers, link redaction and user confirmation."),
])
case_card(s, MARGIN + 2 * (cw + 0.18), 1.8, cw, 3.55, RED, "Meta · open guardrails", "LlamaFirewall for AI agents", [
    ("", "**PromptGuard 2** detects jailbreaks; **AlignmentCheck** audits agent reasoning."),
    ("", "**CodeShield** blocks insecure generated code."),
])
text(s, MARGIN, 5.55, 12.1, 0.3, "WHAT IS PROVING MOST EFFECTIVE", size=12, color=TEAL, bold=True, after=0)
text(s, MARGIN, 5.9, 12.1, 0.8, "**Layers, not one fix:** adversarial training + input/output classifiers + continuous "
     "red teaming + a human check before money moves.", size=15, color=INK)
footer(s, "Sources: Bullwinkel et al. (2025); Google GenAI Security Team (2025); Chennabasappa et al. (2025).", 4)
notes(s, """Industry leaders use three complementary defenses. Microsoft's AI Red Team has attacked more than 100 of its own
generative AI products before release; its lessons are that attackers do not need gradients or deep math to break a
system, that automation helps but human testers remain crucial, and that the work is never complete. That is how companies
test for vulnerabilities. Google defends Gemini with layers, and reports that training the model on adversarial examples
of prompt injection significantly improved its resistance, backed by machine learning classifiers that detect malicious
instructions, removal of suspicious links and asking the user to confirm risky actions. Meta released LlamaFirewall as an
open-source guardrail for AI agents, with PromptGuard 2 to detect jailbreaks, an alignment check that audits the agent's
reasoning, and CodeShield to block insecure code. Across all three, the message is the same: no single defense is
enough, so the most effective strategy is layered.""")

# ---- slide 5: industry comparison chart -----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Industry comparison: attack types vs. defense strategies", "Comparison chart")
COLS = ["Adversarial training", "Red teaming & testing", "Input/output classifiers", "Anomaly detection", "Human verification"]
ROWS = [("Evasion", "Autonomous vehicles", ["●", "●", "◐", "●", "—"]),
        ("Poisoning", "Healthcare, finance", ["◐", "●", "—", "●", "◐"]),
        ("Prompt injection", "Enterprise AI assistants", ["●", "●", "●", "◐", "●"]),
        ("Deepfakes", "Finance", ["—", "◐", "◐", "●", "●"])]
x0, y0, c0w, c1w, cw_, rh = MARGIN, 1.75, 2.15, 2.45, 1.5, 0.74
rect(s, x0, y0, c0w + c1w + 5 * cw_, 0.75, NAVY)
text(s, x0 + 0.15, y0 + 0.15, c0w - 0.2, 0.5, "Attack type", size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, x0 + c0w + 0.1, y0 + 0.15, c1w - 0.2, 0.5, "Industry most hit", size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
for j, c in enumerate(COLS):
    text(s, x0 + c0w + c1w + j * cw_ + 0.05, y0 + 0.08, cw_ - 0.1, 0.6, c, size=12, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, after=0)
SYM = {"●": (TEAL, "Strong"), "◐": (GOLD, "Partial"), "—": (GREY, "Little")}
for i, (atk, ind, vals) in enumerate(ROWS):
    y = y0 + 0.75 + i * rh
    if i % 2 == 0:
        rect(s, x0, y, c0w + c1w + 5 * cw_, rh, LIGHT)
    text(s, x0 + 0.15, y + 0.1, c0w - 0.2, rh - 0.2, atk, size=15, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0 + c0w + 0.1, y + 0.1, c1w - 0.2, rh - 0.2, ind, size=13, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    for j, v in enumerate(vals):
        text(s, x0 + c0w + c1w + j * cw_, y + 0.1, cw_, rh - 0.2, v, size=22, color=SYM[v][0], bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
ly = y0 + 0.75 + 4 * rh + 0.15
text(s, x0, ly, 12.1, 0.4, "**●** strong fit   **◐** partial fit   **—** little effect     "
     "Author's assessment from the cases and company defenses on slides 3–4.", size=12, color=MUTED)
text(s, x0, ly + 0.45, 12.1, 0.6, "**Takeaway:** only red teaming and anomaly detection help against every attack type; deepfakes are "
     "stopped by process (call-back checks), not by retraining.", size=14, color=INK)
footer(s, "Sources: Google GenAI Security Team (2025); Bullwinkel et al. (2025); Chennabasappa et al. (2025); Parisi (2019, Ch. 5).", 5)
notes(s, """This chart compares the four attack types with five defense strategies, and names the industry each attack hits
hardest. Evasion attacks, like the tape on the speed sign, matter most for autonomous vehicles; adversarial training and
robustness testing make the model harder to fool, and anomaly detection can flag an input that does not fit, such as an
85 mile-per-hour reading in a 35 zone. Poisoning threatens healthcare and finance models that retrain on new data; testing
and screening training data for outliers help most. Prompt injection threatens enterprise AI assistants, where every
layer helps, which is why Google and Meta stack them. Deepfakes hit finance; the model cannot be retrained against them,
so anomaly checks on payments and human call-back verification work best. The ratings are my assessment from the cases
and the companies' published defenses.""")

# ---- slide 6: Chapter 5 + governance trend -------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Applying Chapter 5, and the governance trend that requires it", "AI concepts · current trend")
text(s, MARGIN, 1.8, 6.0, 0.3, "CHAPTER 5: ANOMALY DETECTION AS A DEFENSE", size=12, color=TEAL, bold=True, after=0)
text(s, MARGIN, 2.15, 6.0, 0.6, "Model what *normal* looks like, then flag what does not fit.", size=15, color=INK, italic=True)
bullets(s, MARGIN, 2.7, 6.0, 3.9, [
    "**Inputs:** a camera reading of 85 mph on a 35 mph road is out of distribution → slow down, ask the driver.",
    "**Payments:** a large, urgent transfer right after one video call is unusual → hold and call back.",
    "**Training data:** outliers in new data are screened before retraining → stops poisoning.",
    "**AI usage:** an assistant suddenly sending data to an outside link → block it.",
], size=14.5, after=9, marker_color=TEAL)
rect(s, 7.0, 1.8, 5.75, 4.85, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
text(s, 7.3, 2.0, 5.2, 0.3, "TREND: AI GOVERNANCE MAKES ROBUSTNESS LAW", size=12, color=ACCENT, bold=True, after=0)
text(s, 7.3, 2.4, 5.2, 4.2, [
    "The **EU AI Act, Article 15**, requires high-risk AI to be “resilient against attempts by unauthorised third "
    "parties to alter their use, outputs or performance.”",
    "It names data poisoning, model poisoning and adversarial examples, and asks for measures to “prevent, detect, "
    "respond to” them.",
    "**Why it matters:** robustness testing and anomaly detection move from best practice to legal duty.",
], size=14.5, color=INK, after=10)
footer(s, f"Sources: {CH5}; Regulation (EU) 2024/1689 (2024, Art. 15).", 6)
notes(s, """Chapter 5 of the textbook, Network Anomaly Detection with AI, teaches a model of normal behaviour that flags
anything that does not fit. The same idea is a practical defense against adversarial attacks. A speed reading of 85 on a
35 road is out of distribution, so the car should slow down and alert the driver. A large, urgent transfer
right after one video call is unusual, so the payment should be held and confirmed by phone. Outliers in new training data can
be screened before retraining, which blocks poisoning, and an AI assistant that suddenly sends data to an outside link can
be stopped. The current trend is AI governance. Article 15 of the EU AI Act requires high-risk AI systems to resist
attempts to alter their outputs, names data poisoning, model poisoning and adversarial examples, and requires measures to
prevent, detect and respond to them. Together they turn robustness
testing and anomaly detection from good practice into a legal duty.""")

# ---- slide 7: reflection + references --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Reflection and references", "What I learned · APA 7th edition")
rect(s, MARGIN, 1.7, 12.15, 0.95, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, MARGIN + 0.3, 1.82, 11.6, 0.8, [
    "**Reflection.** In Modules 3–5 I built detectors and trusted their accuracy; this module showed that the detector "
    "itself is a target. I want to build AI defenses that are red-teamed before release and backed by anomaly checks.",
], size=14, color=INK)
REFS = [
    "Bullwinkel, B., Minnich, A., Chawla, S., Lopez, G., Pouliot, M., Maxwell, W., de Gruyter, J., Pratt, K., Qi, S., Chikanov, N., Lutz, R., Dheekonda, R. S. R., Jagdagdorj, B.-E., Kim, E., Song, J., Hines, K., Jones, D., Severi, G., Lundeen, R., . . . Russinovich, M. (2025). *Lessons from red teaming 100 generative AI products* (arXiv:2501.07238). arXiv. https://arxiv.org/abs/2501.07238",
    "Chennabasappa, S., Nikolaidis, C., Song, D., Molnar, D., Ding, S., Wan, S., Whitman, S., Deason, L., Doucette, N., Montilla, A., Gampa, A., de Paola, B., Gabi, D., Crnkovich, J., Testud, J.-C., He, K., Chaturvedi, R., Zhou, W., & Saxe, J. (2025). *LlamaFirewall: An open source guardrail system for building secure AI agents* (arXiv:2505.03574). arXiv. https://arxiv.org/abs/2505.03574",
    "Google GenAI Security Team. (2025, June 13). *Mitigating prompt injection attacks with a layered defense strategy*. Google. https://blog.google/security/mitigating-prompt-injection-attacks/",
    "Lalchand, S., Srinivas, V., Maggiore, B., & Henderson, J. (2024, May 29). *Generative AI is expected to magnify the risk of deepfakes and other fraud in banking*. Deloitte Insights. https://www.deloitte.com/us/en/insights/industry/financial-services/deepfake-banking-fraud-risk-on-the-rise.html",
    "Microsoft Security Response Center. (2025, June 11). *CVE-2025-32711: M365 Copilot information disclosure vulnerability*. https://msrc.microsoft.com/update-guide/vulnerability/CVE-2025-32711",
    "Parisi, A. (2019). *Hands-on artificial intelligence for cybersecurity: Implement smart AI systems for preventing cyber attacks and detecting threats and network anomalies*. Packt Publishing.",
    "Povolny, S., & Trivedi, S. (2020, February 19). *Model hacking ADAS to pave safer roads for autonomous vehicles*. McAfee Blogs. https://www.mcafee.com/blogs/other-blogs/mcafee-labs/model-hacking-adas-to-pave-safer-roads-for-autonomous-vehicles/",
    "Regulation (EU) 2024/1689 of the European Parliament and of the Council of 13 June 2024 laying down harmonised rules on artificial intelligence. (2024). *Official Journal of the European Union*, L 2024/1689. http://data.europa.eu/eli/reg/2024/1689/oj",
]
box = text(s, MARGIN, 2.85, SW - 2 * MARGIN, 4.0, REFS, size=11, color=INK, after=4, spacing=1.0)
for p in box.text_frame.paragraphs:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.35))))
    pPr.set("indent", str(-Emu(Inches(0.35))))
footer(s, "Chapter citations: Parisi (2019), Chapter 5 “Network Anomaly Detection with AI” and Chapter 10 “Assessing your AI Arsenal.”", 7)
notes(s, """In Modules 3, 4 and 5 I built an anomaly detector, a fraud model and a phishing filter, and I judged them by accuracy.
This module changed how I see that work: the detector itself is a target, and a model that is 99 percent accurate can be
fooled by a strip of tape or a hidden instruction. As a future security engineer I want to build AI defenses that are
red-teamed before release, backed by anomaly detection, and designed so a human confirms any high-stakes action.""")

prs.save(OUT)
print(f"Wrote {OUT}")
