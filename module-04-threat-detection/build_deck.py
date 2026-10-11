"""Build the Module 4 AI threat detection deck (16:9) with python-pptx.

Usage:  python build_deck.py [executed_notebook.ipynb]
Output: Bandari_Ruthvik_Nath_Module4_AI_Threat_Detection.pptx

Numbers come from the executed notebook's "summary" cell, so the slides match the run that produced them.
"""
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
NOTEBOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Ruthvik_Nath_Bandari_AI_Threat_Detection.ipynb"
OUT = HERE / "Bandari_Ruthvik_Nath_Module4_AI_Threat_Detection.pptx"
DATE = "October 11, 2026"


def load_summary(path):
    nb = json.loads(Path(path).read_text(encoding="utf-8"))
    for cell in nb["cells"]:
        if cell.get("metadata", {}).get("tags") == ["summary"]:
            text = "".join("".join(o.get("text", "")) for o in cell.get("outputs", []) if o.get("name") == "stdout")
            return json.loads(text[text.index("{"):text.rindex("}") + 1])
    sys.exit("Run the notebook to the end first: no summary output found.")


S = load_summary(NOTEBOOK)

# ---- palette ----------------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x1B, 0x2D)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x1F, 0x5F, 0xAD)
GREY = RGBColor(0x9A, 0xA3, 0xAF)
RED = RGBColor(0xC0, 0x39, 0x2B)
TEAL = RGBColor(0x1E, 0x84, 0x6F)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
MARGIN = 0.6


# ---- helpers ----------------------------------------------------------------------------------
def rich(paragraph, text, size, color, bold=False, italic=False):
    """Add runs to a paragraph; **bold** and *italic* markers inside `text` are honoured."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        run = paragraph.add_run()
        is_b, is_i = part.startswith("**"), part.startswith("*") and not part.startswith("**")
        run.text = part.strip("*") if (is_b or is_i) else part
        f = run.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), color
        f.bold = bold or is_b
        f.italic = italic or is_i


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=6, italic=False):
    """Text box. `paras` is a str or list of str (one paragraph each)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p_text in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        rich(p, p_text, size, color, bold, italic)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, after=10, marker_color=ACCENT):
    """Bulleted list with a coloured square marker drawn as a real bullet character."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        p.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.28))))
        pPr.set("indent", str(-Emu(Inches(0.28))))
        bu_clr = pPr.makeelement(qn("a:buClr"), {})
        bu_clr.append(bu_clr.makeelement(qn("a:srgbClr"), {"val": str(marker_color)}))
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        rich(p, item, size, color)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def line(slide, x1, y1, x2, y2, color=RULE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def header(slide, title, kicker):
    rect(slide, 0, 0, 0.18, SH, ACCENT)
    text(slide, MARGIN, 0.42, SW - 2 * MARGIN, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True)
    text(slide, MARGIN, 0.72, SW - 2 * MARGIN, 0.8, title, size=30, color=NAVY, bold=True, spacing=1.0)


def footer(slide, source, n):
    line(slide, MARGIN, SH - 0.55, SW - MARGIN, SH - 0.55, RULE, 0.75)
    text(slide, MARGIN, SH - 0.45, SW - 2 * MARGIN - 0.6, 0.3, source, size=10, color=MUTED)
    text(slide, SW - MARGIN - 0.5, SH - 0.45, 0.5, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = " ".join(body.split())


def card(slide, x, y, w, h, title, tag, rows, color=ACCENT):
    """Organization card: coloured top bar, name, technique tag, then labelled rows."""
    rect(slide, x, y, w, h, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    rect(slide, x, y, w, 0.09, color)
    text(slide, x + 0.25, y + 0.25, w - 0.5, 0.4, title, size=20, color=NAVY, bold=True, after=0)
    text(slide, x + 0.25, y + 0.68, w - 0.5, 0.35, tag, size=12, color=color, bold=True, after=0)
    paras = []
    for label, body in rows:
        paras.append(f"**{label}**")
        paras.append(body)
    box = text(slide, x + 0.25, y + 1.1, w - 0.5, h - 1.25, paras, size=16, color=INK, after=2)
    for i, p in enumerate(box.text_frame.paragraphs):       # gap before each label after the first
        if i and i % 2 == 0:
            p.space_before = Pt(8)
    return box


def bar_chart(slide, x, y, w, h, title, categories, values, colors, number_format, axis_max):
    data = CategoryChartData()
    data.categories = categories
    data.add_series(title, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold = Pt(14), True
    tp.runs[0].font.color.rgb, tp.runs[0].font.name = NAVY, FONT
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = number_format, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.color.rgb, dl.font.name = Pt(16), True, INK, FONT
    for i, c in enumerate(colors):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = axis_max, 0
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.color.rgb, ca.tick_labels.font.name = Pt(12), INK, FONT
    ca.format.line.color.rgb = RULE
    return gf




def grouped_bars(slide, x, y, w, h, categories, series, colors, title):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    r = ch.chart_title.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(14), True, NAVY, FONT
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.name = Pt(12), FONT
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, -10
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0%", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.name = Pt(12), True, FONT
    for s_, c in zip(plot.series, colors):
        s_.format.fill.solid(); s_.format.fill.fore_color.rgb = c
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale, va.visible, va.has_major_gridlines = 1.15, 0, False, False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(13), FONT
    ca.format.line.color.rgb = RULE


def hbars(slide, x, y, w, h, labels, values, color, title):
    data = CategoryChartData()
    data.categories = labels[::-1]
    data.add_series(title, values[::-1])
    ch = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    r = ch.chart_title.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(14), True, NAVY, FONT
    plot = ch.plots[0]
    plot.gap_width = 50
    plot.has_data_labels = True
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = "0.000", False
    plot.data_labels.font.size, plot.data_labels.font.name = Pt(11), FONT
    plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = color
    va = ch.value_axis
    va.visible, va.has_major_gridlines, va.minimum_scale = False, False, 0
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(12), FONT
    ca.format.line.color.rgb = RULE


pct = lambda v, d=1: f"{v * 100:.{d}f}%"
FAM = ["DoS", "Probe", "R2L", "U2R"]
fc = S["family_counts"]
dp = S["dos_profile"]
CITE_C5 = "Parisi (2019, Chapter 5)"

# ---- slide 1: title ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, MARGIN, 1.55, 0.9, 0.08, RED)
text(s, MARGIN, 1.8, 12.0, 1.7, ["Catching Denial-of-Service Attacks", "With an AI Threat Detector"],
     size=44, color=WHITE, bold=True, spacing=1.0, after=0)
text(s, MARGIN, 3.6, 12.0, 1.0, "A Random Forest ensemble trained on NSL-KDD network traffic: what it catches, what it misses, and how a 2025 company should use it",
     size=21, color=RGBColor(0xC9, 0xD3, 0xE0), spacing=1.1)
text(s, MARGIN, 5.25, 12.0, 1.4, ["Ruthvik Nath Bandari",
                                    "AAI6680: AI for Cybersecurity · Northeastern University · Mimoza Dimodugno, PhD", DATE],
     size=16, color=RGBColor(0xC9, 0xD3, 0xE0), after=4)
notes(s, """This deck reports an AI threat detector I built in Google Colab: a Random Forest ensemble trained on the NSL-KDD
network intrusion dataset. I focus on one threat, denial-of-service attacks, show how the detector performs, where it
fails, and how a company in 2025 should deploy it.""")

# ---- slide 2: the threat ------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "The threat: denial-of-service floods that take businesses offline", "The threat and its impact")
bullets(s, MARGIN, 1.85, 7.0, 4.6, [
    "**What it is:** a flood of traffic that overwhelms a server so real customers cannot get in.",
    "**Scale:** Cloudflare blocked **47.1 million** DDoS attacks in 2025, more than double 2024; the largest reached **31.4 Tbps**.",
    "**Financial impact:** an hour of downtime costs **over $300,000** for more than 90% of mid-size and large firms.",
    f"**In the data:** DoS is **{pct(S['dos_share'], 0)}** of NSL-KDD training traffic, the largest attack family.",
], size=17, after=14)
rect(s, 8.2, 1.85, 4.55, 2.75, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
rect(s, 8.2, 1.85, 4.55, 0.09, RED)
text(s, 8.5, 2.15, 4.0, 0.3, "RECENT EXAMPLE", size=12, color=RED, bold=True, after=0)
text(s, 8.5, 2.55, 4.0, 3.7, [
    "**Threema, August 2026.** DDoS attacks left the messaging service unavailable or only partly available for an "
    "evening and the next morning, and also targeted its colocation partner, Nine.",
], size=15, color=INK, after=12)
footer(s, "Sources: Yoachimik et al. (2026); DiDio (2024); Ilascu (2026); Tavallaee et al. (2009).", 2)
notes(s, f"""The one threat I focus on is denial of service. A DoS or distributed DoS attack floods a server with
traffic until legitimate users cannot reach it. Cloudflare reported blocking 47.1 million DDoS attacks in 2025, more than
double the year before, including a record 31.4 terabit-per-second attack. The business impact is direct financial loss: ITIC's
2024 survey found that a single hour of downtime costs more than 300,000 dollars for over 90 percent of mid-size and large
enterprises. A recent example is the August 2026 attack on the Threema messaging service, which was unavailable or only
partly available for an evening and the next morning; the attacks also targeted its colocation partner, Nine. In the NSL-KDD dataset I used, DoS is the largest
attack family, about {pct(S['dos_share'], 0)} of the training connections.""")

# ---- slide 3: approach / Chapter 5 ------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Applying Chapter 5: network anomaly detection with AI", "Approach")
steps = [("1  Load", f"{S['n_raw']:,} labelled connections (NSL-KDD)"),
         ("2  Clean", f"No missing or duplicate rows; constant column dropped; {S['n_features']} features"),
         ("3  Train", "Random Forest: 200 decision trees vote on each connection"),
         ("4  Test", "80/20 split, then the official test file with new attacks")]
bx, by, bw, bh = MARGIN, 1.95, 2.85, 1.55
for i, (t_, d_) in enumerate(steps):
    x = bx + i * (bw + 0.27)
    rect(s, x, by, bw, bh, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, x + 0.2, by + 0.18, bw - 0.4, 0.4, t_, size=17, color=ACCENT, bold=True, after=0)
    text(s, x + 0.2, by + 0.62, bw - 0.4, 0.9, d_, size=13.5, color=INK, after=0)
    if i < 3:
        text(s, x + bw + 0.02, by + 0.5, 0.25, 0.5, "›", size=26, color=GREY, align=PP_ALIGN.CENTER)
text(s, MARGIN, 3.85, 6.0, 0.3, "TWO CHAPTER 5 APPROACHES", size=12, color=TEAL, bold=True, after=0)
bullets(s, MARGIN, 4.25, 6.0, 2.3, [
    "**Supervised classification** (this project): learns labelled attacks; excellent on known threats.",
    "**Anomaly detection** (Modules 1 and 3): learns normal traffic only; flags anything unusual, including new attacks.",
], size=15, after=10, marker_color=TEAL)
text(s, 7.0, 3.85, 5.75, 0.3, "WHY A RANDOM FOREST", size=12, color=ACCENT, bold=True, after=0)
bullets(s, 7.0, 4.25, 5.75, 2.3, [
    "An **ensemble** of many trees is more accurate and stable than one tree, and it shows **which features** drive decisions.",
], size=15, after=10)
footer(s, "Sources: Parisi (2019, Chapter 5); Tavallaee et al. (2009).", 3)
notes(s, f"""The project follows the workflow in Chapter 5 of Parisi's Hands-On Artificial Intelligence for
Cybersecurity, Network Anomaly Detection with AI. I loaded {S['n_raw']:,} labelled connections, checked for missing values,
duplicates and impossible values, dropped one constant column, and one-hot encoded the text fields, giving {S['n_features']}
features. I then trained a Random Forest, an ensemble of 200 decision trees that each see a random sample of the data and
vote. Chapter 5 contrasts two approaches to network intrusion detection. Supervised classification, which I used here,
learns from labelled attacks. Anomaly detection, which I used in Modules 1 and 3, learns only what normal traffic looks
like and flags anything unusual. The results on the next slides show why both are needed.""")

# ---- slide 4: results ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, f"{pct(S['acc'], 2)} accurate on familiar traffic, {pct(S['acc_off'], 1)} with new attacks", "Results · Q1: accuracy")
bullets(s, MARGIN, 1.95, 4.4, 4.5, [
    f"**Held-out split:** {pct(S['acc'], 2)} accuracy, {S['fp']} false alarms in {S['tn'] + S['fp']:,} normal connections.",
    f"**Official test set** ({S['n_novel']} unseen attack types): {pct(S['acc_off'], 1)} accuracy overall.",
    f"**New attack types** caught only {pct(S['novel_rate'], 0)} of the time, against {pct(S['seen_rate'], 0)} for known ones.",
], size=16, after=16)
grouped_bars(s, 5.3, 1.75, 7.45, 4.85, FAM,
             [("Held-out split (known attacks)", [S["split_rate"][f] for f in FAM]),
              ("Official test (includes new attacks)", [S["off_rate"][f] for f in FAM])],
             [ACCENT, RED], "Share of attacks detected, by family")
footer(s, "Source: author's analysis of NSL-KDD (Tavallaee et al., 2009). Held-out U2R bar rests on only 6 connections.", 4)
notes(s, f"""What accuracy did my AI achieve? On a random 20 percent of the training file it reached {pct(S['acc'], 2)}
accuracy, with only {S['fp']} false alarms among {S['tn'] + S['fp']:,} normal connections. That number is flattering,
because the test connections come from the same attacks the model trained on. NSL-KDD's official test file contains
{S['n_novel']} attack types that never appear in training, and there accuracy falls to {pct(S['acc_off'], 1)}. New attack
types were detected only {pct(S['novel_rate'], 0)} of the time, against {pct(S['seen_rate'], 0)} for known ones. The chart
shows the drop for every family; DoS holds up best.""")

# ---- slide 5: why DoS is caught ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Why floods are easy to catch: they look nothing like normal traffic", "Results · Q2: threats detected best")
text(s, MARGIN, 1.85, 5.6, 0.3, "DOS VS NORMAL (MEDIAN PER CONNECTION)", size=12, color=RED, bold=True, after=0)
rows = [("", "DoS", "Normal"),
        ("Connections to same host in 2 s", f"{dp['DoS median']['count']:.0f}", f"{dp['Normal median']['count']:.0f}"),
        ("Share with SYN errors", f"{dp['DoS median']['serror_rate']:.0%}", f"{dp['Normal median']['serror_rate']:.0%}"),
        ("Share to the same service", f"{dp['DoS median']['same_srv_rate']:.0%}", f"{dp['Normal median']['same_srv_rate']:.0%}"),
        ("Bytes sent by source", f"{dp['DoS median']['src_bytes']:.0f}", f"{dp['Normal median']['src_bytes']:.0f}")]
for r_, row in enumerate(rows):
    yy = 2.3 + r_ * 0.62
    for c_, (cell, cx, cw_) in enumerate(zip(row, (MARGIN, MARGIN + 3.4, MARGIN + 4.6), (3.3, 1.1, 1.1))):
        text(s, cx, yy, cw_, 0.5, cell, size=15, color=(RED if c_ == 1 and r_ else MUTED if r_ == 0 or c_ == 0 else INK),
             bold=(r_ == 0 or c_ == 1), align=PP_ALIGN.LEFT if c_ == 0 else PP_ALIGN.RIGHT)
    if r_:
        line(s, MARGIN, yy + 0.52, MARGIN + 5.7, yy + 0.52, RULE, 0.75)
text(s, MARGIN, 5.55, 5.7, 0.9, f"DoS was detected {pct(S['off_rate']['DoS'])} of the time on the official test set, the best of any family.",
     size=15, color=INK, bold=True)
tf = S["top_features"]
names = list(tf)[:8]
hbars(s, 6.75, 1.75, 6.0, 4.4, names, [tf[n] for n in names], TEAL, "What the Random Forest relies on (importance)")
text(s, 6.75, 6.2, 6.0, 0.4, "Importance across all attacks: byte counts, connection status and traffic patterns separate floods.",
     size=12, color=MUTED, italic=True)
footer(s, f"Sources: author's analysis; {CITE_C5}.", 5)
notes(s, f"""What threats does it detect best? Denial of service. A flood is statistically loud: the median DoS
connection is one of {dp['DoS median']['count']:.0f} connections to the same host in the previous two seconds, against
{dp['Normal median']['count']:.0f} for normal traffic, and nearly all of them end in SYN errors because the attacker never
completes the handshake. The Random Forest picks this up mainly through byte counts, the connection-status flag and connection-pattern
features such as count and same-service rate, shown on the right.
That is why DoS stayed the best-detected family on the official test set, at {pct(S['off_rate']['DoS'])}.""")

# ---- slide 6: limits / top threats ----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "The real danger: the attacks the detector misses", "Top threats and limits")
tm = S["top_missed"][:6]
text(s, MARGIN, 1.85, 6.4, 0.3, "MOST-MISSED ATTACK TYPES (OFFICIAL TEST SET)", size=12, color=RED, bold=True, after=0)
for r_, t_ in enumerate([{"label": "Attack", "family": "Family", "missed": "Missed", "new_type": None}] + tm):
    yy = 2.25 + r_ * 0.55
    hdr = r_ == 0
    cells = [t_["label"], t_["family"], f"{t_['missed']:,}" if not hdr else "Missed",
             "" if hdr else ("new" if t_["new_type"] else "known")]
    for cell, cx, cw_, al in zip(cells, (MARGIN, MARGIN + 2.5, MARGIN + 3.7, MARGIN + 5.0), (2.4, 1.1, 1.2, 1.3),
                                 (PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.LEFT)):
        text(s, cx, yy, cw_, 0.45, cell, size=14.5, color=MUTED if hdr else INK, bold=hdr, align=al)
    if r_:
        line(s, MARGIN, yy + 0.47, MARGIN + 6.3, yy + 0.47, RULE, 0.75)
text(s, 7.35, 1.85, 5.4, 0.3, "WHAT THIS MEANS", size=12, color=ACCENT, bold=True, after=0)
bullets(s, 7.35, 2.25, 5.4, 4.3, [
    "**New DoS variants** (processtable, apache2) slip past: the model never saw them in training.",
    f"**Quieter attacks** such as password guessing are missed too: only {pct(S['off_rate']['R2L'])} of R2L attacks were caught.",
    "**Mitigation (Chapter 5):** add anomaly detection, which flags unusual traffic it was never shown.",
], size=15.5, after=14)
footer(s, f"Sources: author's analysis; {CITE_C5}.", 6)
notes(s, f"""Which network activities are most dangerous? The ones the model misses. For our threat, the danger is new
DoS variants: processtable and apache2 never appear in training, and most of them slipped through. Quieter attacks were
missed too; the most-missed type overall was password guessing, and the R2L family was caught only
{pct(S['off_rate']['R2L'])} of the time. A supervised model only recognises what it was trained on. Chapter 5 suggests the mitigation: combine the classifier with
anomaly detection, which models normal traffic and flags anything that departs from it, including attacks with no label.""")

# ---- slide 7: 2025 company + trend --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Trend: autonomous DDoS mitigation, and where this detector fits", "Q3: value to a company · current trend")
rect(s, MARGIN, 1.85, 6.0, 2.45, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
text(s, MARGIN + 0.3, 2.05, 5.4, 0.3, "TREND: AUTONOMOUS DDOS MITIGATION", size=12, color=TEAL, bold=True, after=0)
text(s, MARGIN + 0.3, 2.45, 5.4, 3.9, [
    "Cloudflare blocked about **5,343 network-layer DDoS attacks per hour** in the first half of 2026. Its systems act "
    "without human intervention, because attacks end before an alert would reach an analyst.",
], size=15.5, color=INK, after=14)
text(s, 7.0, 1.9, 5.75, 0.3, "HOW A COMPANY SHOULD USE THIS DETECTOR", size=12, color=ACCENT, bold=True, after=0)
bullets(s, 7.0, 2.3, 5.75, 4.2, [
    "**Automate the confident cases:** block clear floods in real time.",
    f"**Route uncertain alerts to analysts:** on new traffic, {pct(S['fpr_off'], 1)} of normal connections were wrongly flagged.",
    "**Add anomaly detection and retrain** as new attacks appear.",
], size=15.5, after=14)
footer(s, "Sources: Cloudforce One (2026); author's analysis.", 7)
notes(s, f"""How could this help a real company in 2025? The current trend in DDoS defence is autonomous mitigation.
Cloudflare reports blocking roughly 5,343 network-layer attacks per hour in the first half of 2026, with systems that
detect and mitigate without human intervention, because an attack is usually over before an alert could reach an analyst. My
detector fits that model for floods: DoS is its strongest family, so clear cases can be blocked automatically. But on the
official test set it also flagged {pct(S['fpr_off'], 1)} of normal connections as attacks, and it missed most new attack
types. So the responsible design automates only high-confidence decisions, sends uncertain ones to analysts, adds anomaly
detection for the unknown, and retrains regularly.""")

# ---- slide 8: reflection ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Reflection: accuracy is not the same as protection", "What I learned")
rect(s, MARGIN, 1.85, 12.15, 3.3, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
text(s, MARGIN + 0.4, 2.15, 11.4, 4.1, [
    f"My first result, **{pct(S['acc'], 2)}**, looked finished. The official test set (**{pct(S['acc_off'], 1)}**, and only "
    f"**{pct(S['novel_rate'], 0)}** of new attacks caught) showed the model knew the attacks it had seen, not attacks in general.",
    "My Module 1 and 3 anomaly detectors had the opposite profile: supervised and anomaly-based detection cover each other's blind spots.",
    "As a future security engineer, I want to build detectors that are judged on the threats they miss, and that let "
    "people make the calls a model is not sure about.",
], size=18, color=INK, after=16)
footer(s, "Reflection draws on my Module 1, Module 3 and Module 4 labs.", 8)
notes(s, """My first result looked like a finished project, but the official test set showed that the model only knew
the attacks it had already seen. My anomaly detectors in Modules 1 and 3 had the opposite strengths, catching new behaviour
but with more false alarms. The lesson I take into my career is that a detector should be judged by what it misses, and
that the best systems combine automated speed with human judgement on uncertain cases.""")

# ---- slide 9: references ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "References", "APA 7th edition")
REFS = [
    "Cloudforce One. (2026, August 11). *Cloudflare DDoS threat report H1 2026: 1 Tbps attacks soar as DNS floods and geopolitical tensions drive a new wave*. Cloudflare Blog. https://blog.cloudflare.com/ddos-threat-report-2026-h1/",
    "DiDio, L. (2024, April). *ITIC 2024 hourly cost of downtime survey results, part 2*. Information Technology Intelligence Consulting; Calyptix Security. https://calyptix.com/wp-content/uploads/ITIC-2024-Hourly-Cost-of-Downtime-Survey-Results-Part-2.pdf",
    "Ilascu, I. (2026, August 16). Large-scale DDoS attacks disrupted Threema secure messaging service. *BleepingComputer*. https://www.bleepingcomputer.com/news/security/large-scale-ddos-attacks-disrupted-threema-secure-messaging-service/",
    "Parisi, A. (2019). *Hands-on artificial intelligence for cybersecurity: Implement smart AI systems for preventing cyber attacks and detecting threats and network anomalies*. Packt Publishing.",
    "Tavallaee, M., Bagheri, E., Lu, W., & Ghorbani, A. A. (2009). A detailed analysis of the KDD CUP 99 data set. In *Proceedings of the 2009 IEEE Symposium on Computational Intelligence for Security and Defense Applications* (pp. 1–6). IEEE. https://doi.org/10.1109/CISDA.2009.5356528",
    "Yoachimik, O., Pacheco, J., & Cloudforce One. (2026, February 5). *2025 Q4 DDoS threat report: A record-setting 31.4 Tbps attack caps a year of massive DDoS assaults*. Cloudflare Blog. https://blog.cloudflare.com/ddos-threat-report-2025-q4/",
]
box = text(s, MARGIN, 1.8, SW - 2 * MARGIN, 5.0, REFS, size=12.5, color=INK, after=8, spacing=1.05)
for p in box.text_frame.paragraphs:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.45))))
    pPr.set("indent", str(-Emu(Inches(0.45))))
footer(s, "Notebook: Ruthvik_Nath_Bandari_AI_Threat_Detection.ipynb (Google Colab).", 9)
notes(s, "References in APA 7th edition, including Chapter 5 of Parisi (2019) and recent Cloudflare threat reports.")

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides")
