"""Build the Module 5 phishing detective deck (16:9) with python-pptx.

Usage:  python build_deck.py [executed_notebook.ipynb] [notebook_share_link]
Output: Bandari_Ruthvik_Nath_Module5_Phishing_Detective.pptx

Numbers and charts come from the executed notebook (its "summary" cell and its saved figure outputs), so the slides
match the run that produced them.
"""
import base64
import io
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
NOTEBOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Ruthvik_Nath_Bandari_Phishing_Detective.ipynb"
LINK = sys.argv[2] if len(sys.argv) > 2 else "[OneDrive link to be added]"
OUT = HERE / "Bandari_Ruthvik_Nath_Module5_Phishing_Detective.pptx"
DATE = "October 18, 2026"

NB = json.loads(NOTEBOOK.read_text(encoding="utf-8"))


def load_summary():
    for cell in NB["cells"]:
        if cell.get("metadata", {}).get("tags") == ["summary"]:
            text = "".join("".join(o.get("text", "")) for o in cell.get("outputs", []) if o.get("name") == "stdout")
            return json.loads(text[text.index("{"):text.rindex("}") + 1])
    sys.exit("Run the notebook to the end first: no summary output found.")


def figure(name):
    """PNG output of the code cell that saves figures/<name>.png."""
    for cell in NB["cells"]:
        if cell["cell_type"] == "code" and f"{name}.png" in "".join(cell["source"]):
            for o in cell.get("outputs", []):
                if "image/png" in o.get("data", {}):
                    return io.BytesIO(base64.b64decode(o["data"]["image/png"]))
    sys.exit(f"No image output for {name}")


S = load_summary()
LR, RF = S["results"]["Logistic regression"], S["results"]["Random forest"]
CM = S["lr_cm"]

# ---- palette ----------------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x1B, 0x2D)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD5, 0xDA, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x1F, 0x5F, 0xAD)
GREY = RGBColor(0x9A, 0xA3, 0xAF)
RED = RGBColor(0xC0, 0x39, 0x2B)
TEAL = RGBColor(0x1E, 0x84, 0x6F)
SOFT = RGBColor(0xC9, 0xD3, 0xE0)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
MARGIN = 0.6


# ---- helpers ----------------------------------------------------------------------------------
def rich(paragraph, text, size, color, bold=False, italic=False):
    """Add runs to a paragraph; **bold** and *italic* markers inside `text` are honoured."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        run = paragraph.add_run()
        is_b, is_i = part.startswith("**"), part.startswith("*") and not part.startswith("**")
        run.text = part.strip("*") if (is_b or is_i) else part
        f = run.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), color
        f.bold = bold or is_b
        f.italic = italic or is_i


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=6, italic=False):
    """Text box. `paras` is a str or list of str (one paragraph each)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p_text in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        rich(p, p_text, size, color, bold, italic)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, after=10, marker_color=ACCENT):
    """Bulleted list with a coloured square marker drawn as a real bullet character."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        p.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.28))))
        pPr.set("indent", str(-Emu(Inches(0.28))))
        bu_clr = pPr.makeelement(qn("a:buClr"), {})
        bu_clr.append(bu_clr.makeelement(qn("a:srgbClr"), {"val": str(marker_color)}))
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        rich(p, item, size, color)
    return box


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def line(slide, x1, y1, x2, y2, color=RULE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def header(slide, title, kicker):
    rect(slide, 0, 0, 0.18, SH, ACCENT)
    text(slide, MARGIN, 0.42, SW - 2 * MARGIN, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True)
    text(slide, MARGIN, 0.72, SW - 2 * MARGIN, 0.8, title, size=30, color=NAVY, bold=True, spacing=1.0)


def footer(slide, source, n):
    line(slide, MARGIN, SH - 0.55, SW - MARGIN, SH - 0.55, RULE, 0.75)
    text(slide, MARGIN, SH - 0.45, SW - 2 * MARGIN - 0.6, 0.3, source, size=10, color=MUTED)
    text(slide, SW - MARGIN - 0.5, SH - 0.45, 0.5, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = " ".join(body.split())


def picture(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(figure(name), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def stat(slide, x, y, w, value, label, color=ACCENT):
    rect(slide, x, y, w, 1.35, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(slide, x, y, 0.08, 1.35, color)
    text(slide, x + 0.3, y + 0.14, w - 0.45, 0.6, value, size=30, color=color, bold=True, after=0)
    text(slide, x + 0.3, y + 0.78, w - 0.45, 0.55, label, size=13, color=INK, after=0, spacing=1.0)


pct = lambda v, d=1: f"{v * 100:.{d}f}%"
CMP = S["compare"]
ratio = lambda k: CMP[k]["Phishing Email"] / CMP[k]["Safe Email"]
CH3, CH5 = "Parisi (2019, Ch. 3)", "Parisi (2019, Ch. 5)"
words = S["top_phish_words"]

# ---- slide 1: title ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, MARGIN, 1.4, 0.9, 0.08, RED)
text(s, MARGIN, 1.65, 12.0, 1.7, ["Phishing Detective", "An AI System That Scores Every Email"],
     size=44, color=WHITE, bold=True, spacing=1.0, after=0)
text(s, MARGIN, 3.45, 12.0, 1.0, f"Logistic regression trained on {S['n_train']:,} labelled emails, tested on {S['n_test']:,} more, "
     "with a safety score for every message", size=21, color=SOFT, spacing=1.1)
text(s, MARGIN, 4.75, 12.0, 1.4, ["Ruthvik Nath Bandari",
                                    "AAI6680: AI for Cybersecurity · Northeastern University · Mimoza Dimodugno, PhD", DATE],
     size=16, color=SOFT, after=4)
text(s, MARGIN, 6.3, 12.0, 0.6, f"**Google Colab notebook (OneDrive):** {LINK}", size=13, color=SOFT)
notes(s, """This deck presents a phishing detector I built in Google Colab. It reads an email, decides whether it is
phishing, and gives it a safety score from 0 to 100, like the filters in Gmail and Outlook. The full notebook, with the
working model and all charts, is at the OneDrive link on this slide.""")

# ---- slide 2: the threat ------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "The threat: phishing emails that steal credentials and money", "The threat and its impact")
stat(s, MARGIN, 1.85, 3.9, "191,561", "phishing/spoofing complaints in 2025: the most-reported crime type", RED)
stat(s, MARGIN + 4.12, 1.85, 3.9, "$3.05 billion", "lost to business email compromise (BEC) in 2025", RED)
stat(s, MARGIN + 8.24, 1.85, 3.9, "$20.9 billion", "total cybercrime losses reported to the FBI in 2025, up 26%", ACCENT)
bullets(s, MARGIN, 3.6, 12.1, 3.0, [
    "**What it is:** an email posing as a bank, a courier or your boss that asks you to click, log in or pay.",
    "**Impact:** direct **financial loss** (stolen accounts, wire fraud, gift cards) and a way in for ransomware.",
    "**Why AI:** nobody can inspect every message; an AI filter scores each one before it reaches the inbox.",
], size=17, after=12)
footer(s, "Source: Federal Bureau of Investigation (FBI, 2026).", 2)
notes(s, """The one threat this project focuses on is phishing. A phishing email impersonates someone you trust and tries
to get you to click a link, hand over a password, or send money. According to the FBI's 2025 IC3 Annual Report,
phishing and spoofing was the most-reported crime type, with 191,561 complaints, and business email compromise, where a
fake email from an executive or supplier redirects a payment, cost victims just over 3 billion dollars. Total reported
cybercrime losses reached 20.9 billion dollars, up 26 percent on 2024. The impact is direct financial loss, and phishing is
also the usual first step of larger attacks.""")

# ---- slide 3: approach / Chapter concepts --------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Chapter 5 concepts: machine learning + anomaly detection", "Approach · AI concepts")
steps = [("1  Load", f"{S['n_raw']:,} Kaggle emails: safe, or phishing and spam"),
         ("2  Clean", f"Drop empty and duplicate emails, strip HTML → {S['n_clean']:,}"),
         ("3  Features", "Words (TF-IDF), urgent words, personal-info requests, length"),
         ("4  Train & test", f"80/20 split: {S['n_train']:,} / {S['n_test']:,}"),
         ("5  Score", "Safety score 0–100 for any new email")]
bx, by, bw, bh = MARGIN, 1.9, 2.25, 1.75
for i, (t_, d_) in enumerate(steps):
    x = bx + i * (bw + 0.22)
    rect(s, x, by, bw, bh, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, x + 0.18, by + 0.16, bw - 0.36, 0.4, t_, size=16, color=ACCENT, bold=True, after=0)
    text(s, x + 0.18, by + 0.58, bw - 0.36, 1.1, d_, size=13, color=INK, after=0)
    if i < 4:
        text(s, x + bw + 0.0, by + 0.6, 0.22, 0.5, "›", size=24, color=GREY, align=PP_ALIGN.CENTER)
text(s, MARGIN, 4.0, 6.0, 0.3, "MACHINE LEARNING: SUPERVISED CLASSIFICATION", size=12, color=ACCENT, bold=True, after=0)
bullets(s, MARGIN, 4.4, 6.0, 2.4, [
    "**Logistic regression** (main model) gives every word a weight, showing *why* an email is flagged.",
    "**Random forest** (300 decision trees) for comparison.",
], size=15, after=10)
text(s, 7.0, 4.0, 5.75, 0.3, "ANOMALY DETECTION: THE SAFETY NET", size=12, color=TEAL, bold=True, after=0)
bullets(s, 7.0, 4.4, 5.75, 2.4, [
    "A classifier knows only scams it has seen; anomaly detection flags **unusual** behaviour.",
    "*Example:* a first-time sender asking for a wire transfer → hold for review.",
], size=15, after=10, marker_color=TEAL)
footer(s, f"Sources: {CH3}; {CH5}; Cyber Cop (2023).", 3)
notes(s, f"""The pipeline follows the machine learning workflow in Parisi's textbook. Chapter 3, Ham or Spam, applies
supervised classifiers such as logistic regression and decision trees to email. I loaded {S['n_raw']:,} labelled emails
from Kaggle, removed empty and duplicate messages and HTML, and kept {S['n_clean']:,}. Each email became TF-IDF word
weights plus hand-made features for urgent words, requests for personal information, and length and structure. I trained
logistic regression, which can explain its decisions word by word, and compared it with a random forest. Chapter 5, on
anomaly detection, adds the idea I use for mitigation: a supervised model recognises only known patterns, so an anomaly
detector that flags unusual behaviour is a useful second layer.""")

# ---- slide 4: phishing vs normal -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "What phishing looks like compared with normal email", "Data · phishing vs normal characteristics")
picture(s, "fig_compare", MARGIN, 1.7, w=7.9)
text(s, 8.85, 1.85, 3.9, 0.3, "PER EMAIL, PHISHING HAS", size=12, color=RED, bold=True, after=0)
bullets(s, 8.85, 2.25, 3.9, 4.4, [
    f"**{ratio('exclamations'):.1f}×** as many exclamation marks",
    f"**{ratio('money_signs'):.1f}×** as many money or prize terms",
    f"**{ratio('personal_info_requests'):.1f}×** as many requests for passwords, card or account numbers",
    f"**{ratio('urgent_words'):.1f}×** as many urgent words",
    "Links are **not** a giveaway: safe emails had slightly more.",
], size=15, after=10, marker_color=RED)
footer(s, f"Source: author's analysis of {S['n_clean']:,} cleaned emails from Cyber Cop (2023).", 4)
notes(s, f"""Before training, I compared the two classes. Phishing emails had {ratio('exclamations'):.1f} times as many
exclamation marks, {ratio('money_signs'):.1f} times as many money or prize terms, {ratio('personal_info_requests'):.1f}
times as many requests for personal information and {ratio('urgent_words'):.1f} times as many urgent words. Links on their
own were not a giveaway: safe emails, many of them mailing-list posts, had slightly more. No single feature separates the
classes, which is why the model combines them with the words themselves.""")

# ---- slide 5: accuracy ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, f"{pct(LR['accuracy'])} accurate on {S['n_test']:,} emails it had never seen", "Results · model accuracy")
picture(s, "fig_accuracy", MARGIN, 1.75, w=8.3)
stat(s, 9.25, 1.85, 3.5, pct(LR["recall"]), f"of phishing caught ({CM['tp']:,} of {CM['tp'] + CM['fn']:,})", TEAL)
stat(s, 9.25, 3.4, 3.5, f"{CM['fp']} of {CM['tn'] + CM['fp']:,}", "safe emails wrongly flagged", RED)
text(s, 9.25, 5.0, 3.5, 1.5, f"Logistic regression beat the random forest ({pct(RF['accuracy'])}), and its word weights show why it flags an email.",
     size=14, color=INK)
footer(s, "Source: author's analysis; test set = 20% of emails held out before training.", 5)
notes(s, f"""How accurate is the detector? On {S['n_test']:,} test emails it had never seen, logistic regression was
{pct(LR['accuracy'])} accurate, with precision {LR['precision']:.3f} and recall {LR['recall']:.3f}. It caught {CM['tp']:,}
of {CM['tp'] + CM['fn']:,} phishing emails, missing {CM['fn']}, and wrongly flagged {CM['fp']} of {CM['tn'] + CM['fp']:,}
safe emails. The random forest reached {pct(RF['accuracy'])}. Duplicates were removed before the split, so the score is not
inflated by the same email appearing in training and test.""")

# ---- slide 6: suspicious words -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "The most suspicious words and phrases the AI learned", "Results · dangerous patterns")
picture(s, "fig_words", MARGIN, 1.85, w=7.6)
text(s, 8.45, 1.85, 4.3, 0.3, "WHAT THE MODEL TREATS AS DANGEROUS", size=12, color=RED, bold=True, after=0)
bullets(s, 8.45, 2.25, 4.3, 4.4, [
    f"**Top phishing words:** {', '.join(f'*{w}*' for w in words[:5])}: links, calls to act and offers.",
    "**Exclamation marks** and **money terms** also push toward phishing.",
    "**Caveat:** names and places (*vince*, *houston*) mark the safe archive, so accuracy holds for this dataset only.",
], size=15, after=12, marker_color=RED)
footer(s, "Source: author's analysis; weights from the logistic regression.", 6)
notes(s, f"""Which patterns does the AI find most dangerous? The logistic regression weights show it: calls to act such as
{', '.join(words[:5])}, and money and offer language. On the safe side are conversational words such as thanks and wrote,
names and university addresses. I removed numbers and the archive names Enron and SpamAssassin from the vocabulary,
because the two classes were collected in different years, but people's names and places from the safe archive still
appear. That is why I treat the accuracy as valid for this dataset, not for any inbox.""")

# ---- slide 7: demo -------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Demo: email safety scores for five emails I wrote", "Security demo")
picture(s, "fig_demo", MARGIN, 1.7, w=8.0)
text(s, 8.95, 1.85, 3.8, 0.3, "SCORE = 100 × P(SAFE)", size=12, color=ACCENT, bold=True, after=0)
bullets(s, 8.95, 2.25, 3.8, 2.0, ["**70–100** safe", "**40–69** be careful", "**0–39** likely phishing"],
        size=15, after=6)
sus = [d for d in S["demo"] if d["written as"] == "Suspicious"]
text(s, 8.95, 3.85, 3.8, 2.8, [
    f"All {len(S['demo'])} verdicts matched my own label. The CEO-style *Quick favour* email has no link, yet scored "
    f"**{sus[-1]['safety score']}**; top signals: {sus[-1]['top phishing signals']}.",
], size=14, color=INK)
footer(s, "Source: author's analysis; demo emails written for this test, not taken from the dataset.", 7)
notes(s, f"""To test the system like a user would, I wrote five emails: two normal workplace messages and three
scams, an account-suspension warning, a gift-card prize and a CEO-style wire-transfer request. The safety score is 100
times the model's probability that the email is safe. The safe emails scored {' and '.join(str(d['safety score']) for d in S['demo'] if d['written as'] == 'Safe')},
and the scams scored {', '.join(str(d['safety score']) for d in sus)}. The wire-transfer email has no link, so the model
flagged it from its wording and length: {sus[-1]['top phishing signals']}.""")

# ---- slide 8: current trend ------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Trend: AI-written phishing, and how to defend against it", "Current trend · mitigation")
rect(s, MARGIN, 1.85, 6.0, 3.6, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
text(s, MARGIN + 0.3, 2.05, 5.4, 0.3, "TREND: GENERATIVE AI IN PHISHING", size=12, color=RED, bold=True, after=0)
text(s, MARGIN + 0.3, 2.45, 5.4, 3.9, [
    "In 2025, **22,364** complaints to the FBI mentioned AI, with **$893 million** in losses.",
    "Chat generators “can quickly create official-sounding emails mimicking a company’s CEO” (FBI, 2026, p. 39); "
    "BEC scams involving AI cost businesses **over $30 million**.",
    "The tell-tale spelling mistakes disappear.",
], size=15.5, color=INK, after=12)
text(s, 7.0, 1.9, 5.75, 0.3, "HOW TO MITIGATE IT", size=12, color=TEAL, bold=True, after=0)
bullets(s, 7.0, 2.3, 5.75, 4.2, [
    f"**Score wording, not spelling:** the link-free *Quick favour* scam still scored {sus[-1]['safety score']} from words such as *account* and *bank*.",
    "**Add anomaly detection** (Chapter 5): hold first-time senders who ask for payment.",
    "**Verify out of band:** confirm any money request by phone.",
    "**Retrain regularly** on recent, AI-written phishing.",
], size=15, after=11, marker_color=TEAL)
footer(s, f"Sources: FBI (2026); {CH5}.", 8)
notes(s, """The current trend I chose is attackers using generative AI. The FBI's 2025 report has a new section on AI:
22,364 complaints mentioned it, with 893 million dollars in losses, and it warns that chat generators can quickly produce
official-sounding emails that mimic a CEO. That removes the spelling mistakes people were taught to look for. My
detector still has a role because it scores wording rather than spelling, and the wire-transfer demo shows it can catch a
message with no link at all. The mitigation is layered: the content model, anomaly
detection for unusual senders and requests as in Chapter 5, out-of-band verification of payments, and regular retraining
on recent AI-written phishing.""")

# ---- slide 9: report & reflection -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Security report and reflection", "Report · what I learned")
qa = [("How accurate is it?", f"{pct(LR['accuracy'])} on unseen emails from this dataset: {pct(LR['recall'])} of phishing caught, {CM['fp']} false alarms in {CM['tn'] + CM['fp']:,} safe emails."),
      ("What looks most dangerous?", f"Words such as {', '.join(words[:3])} and money, plus exclamation marks; phishing also asked for passwords or card numbers {ratio('personal_info_requests'):.1f}× as often."),
      ("How could it protect my family?", "Score every email in the inbox and warn before anyone clicks or pays, especially on account-suspension and gift-card scams.")]
for i, (q, a) in enumerate(qa):
    y = 1.85 + i * 1.05
    text(s, MARGIN, y, 3.4, 0.9, q, size=15, color=ACCENT, bold=True, after=0)
    text(s, MARGIN + 3.5, y, 8.6, 0.9, a, size=15, color=INK, after=0)
    line(s, MARGIN, y + 0.92, SW - MARGIN, y + 0.92, RULE, 0.75)
rect(s, MARGIN, 5.1, 12.15, 1.6, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, MARGIN + 0.3, 5.25, 11.6, 1.4, [
    "**Reflection.** Modules 3 and 4 applied AI to traffic and transactions; here the same workflow worked on language. "
    "I want to build email defences that pair models like this with anomaly detection and plain warnings my family can act on.",
], size=15, color=INK)
footer(s, "Full security report (100–150 words) in the notebook's Security report section.", 9)
notes(s, """The notebook's security report answers the three questions in under 150 words. The detector is highly
accurate on this dataset, it weights calls to act, money language and exclamation marks most heavily, and it
could protect my family by scoring every email before anyone clicks. It still misses some phishing, so it supports
rather than replaces checking with the sender. In Modules 3 and 4 I applied AI to network traffic and transactions; this
project showed the same workflow works on language. I want to build email defences that combine content models with
anomaly detection and explain their warnings in plain words.""")

# ---- slide 10: references ----------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "References", "APA 7th edition")
REFS = [
    "Cyber Cop. (2023). *Phishing email detection* [Data set]. Kaggle. https://www.kaggle.com/datasets/subhajournal/phishingemails",
    "Federal Bureau of Investigation. (2026). *2025 IC3 annual report*. Internet Crime Complaint Center. "
    "https://www.ic3.gov/AnnualReport/Reports/2025_IC3Report.pdf",
    "Parisi, A. (2019). *Hands-on artificial intelligence for cybersecurity: Implement smart AI systems for preventing "
    "cyber attacks and detecting threats and network anomalies*. Packt Publishing.",
]
box = text(s, MARGIN, 1.8, SW - 2 * MARGIN, 3.5, REFS, size=14, color=INK, after=12, spacing=1.05)
for p in box.text_frame.paragraphs:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.45))))
    pPr.set("indent", str(-Emu(Inches(0.45))))
text(s, MARGIN, 5.2, SW - 2 * MARGIN, 0.9, [f"**Notebook:** Ruthvik_Nath_Bandari_Phishing_Detective.ipynb (Google Colab) · {LINK}"],
     size=13, color=MUTED)
footer(s, "Chapter citations: Parisi (2019), Chapter 3 “Ham or Spam?” and Chapter 5 “Network Anomaly Detection with AI.”", 10)
notes(s, "References in APA 7th edition: the course textbook (Chapters 3 and 5), the 2025 FBI IC3 report and the Kaggle dataset.")

prs.save(OUT)
print(f"Wrote {OUT}")
